"""Builds the lookup template for the Dephub ship registry (https://kapal.dephub.go.id/ditkapel_service/data_kapal/) from the vessel workbook.

  python -I docs/scripts/build-vessel-registry-template.py [vessel_workbook.xlsx] [sheet] [output.xlsx]
  (defaults: docs/Vessel Cleanup (Jovin, Klip, SAP) v7.xlsx / Vessel_Prefix -> docs/Vessel Registry Lookup (Dephub).xlsx)

The registry is looked up BY HAND: its robots.txt is `Disallow: /` for every robot, so nothing here fetches from it. The template lists each distinct vessel
once, with the text to type into the search box (the name without its BG. / TB. / MT. prefix and voyage number), the vessels used by KLIP transactions first.
Whoever searches pastes the result rows into the sheet `Hasil`; match-vessel-registry.py then matches them back and writes the English columns.
Needs openpyxl.
"""
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from vessel_common import key, search_name  # noqa: E402

SRC = sys.argv[1] if len(sys.argv) > 1 else 'D:/Project/Klip/docs/Vessel Cleanup (Jovin, Klip, SAP) v7.xlsx'
SHEET = sys.argv[2] if len(sys.argv) > 2 else 'Vessel_Prefix'
OUT = sys.argv[3] if len(sys.argv) > 3 else 'D:/Project/Klip/docs/Vessel Registry Lookup (Dephub).xlsx'

# the columns of the registry's result table, in its own order and wording, so a row can be pasted as it is
SITE_COLUMNS = ['No', 'Nama Kapal', 'Eks Nama Kapal', 'Call Sign', 'Jenis Kapal', 'Nama Pemilik', 'TPK', 'Panjang', 'Lebar', 'Dalam', 'LOA', 'GT', 'Isi Bersih', 'IMO', 'Tahun', 'Mesin']

# two real lookups, done by hand on 2026-10-13: they are genuine rows, and they show the paste format
EXAMPLE_ROWS = [
    ['SOLID 10', 1, 'SOLID 10', '', '', 'Barge', 'PT. PELAYARAN ARGO MITRA SUKSES', '2016 PPm No. 4203/L', 52.67, 19.51, 4.27, 54.86, 1112, 334, '', 2015, ''],
    ['ANAK LAUT 289', 1, 'ANAK LAUT 289', 'DOLPHIN', '', 'Barge', 'PT. BUANA KENCANA MARITIME', '2012 HHa No. 2861/L', 79.01, 21.34, 4.88, 82.30, 2146, 1553, '', 2003, ''],
]

wb_in = openpyxl.load_workbook(SRC, read_only=True)
data = [list(r) for r in wb_in[SHEET].iter_rows(values_only=True)]
header, rows = data[0], data[1:]
col = {h: i for i, h in enumerate(header)}
need = ['Pair ID', 'Nama Vessel', 'Transaksi Klip', 'Vessel Type', 'Vessel Code SAP', 'Vessel Owner (Company Name)', 'Vessel LOA', 'Vessel Draft', 'Vessel Registration Year']
missing = [h for h in need if h not in col]
if missing:
    sys.exit(f'sheet {SHEET} lacks the columns: {missing}')


def cell(r, name):
    v = r[col[name]]
    return v if v not in ('',) else None


vessels = {}
for r in rows:
    name = cell(r, 'Nama Vessel')
    if not name:
        continue
    k = key(name)
    v = vessels.setdefault(k, {'name': name, 'pairs': [], 'tx': 0, 'type': None, 'code': None, 'owner': None, 'loa': None, 'draft': None, 'year': None})
    if cell(r, 'Pair ID') and cell(r, 'Pair ID') not in v['pairs']:
        v['pairs'].append(cell(r, 'Pair ID'))
    v['tx'] = max(v['tx'], cell(r, 'Transaksi Klip') or 0)
    for f, c in (('type', 'Vessel Type'), ('code', 'Vessel Code SAP'), ('owner', 'Vessel Owner (Company Name)'), ('loa', 'Vessel LOA'), ('draft', 'Vessel Draft'), ('year', 'Vessel Registration Year')):
        if v[f] is None and cell(r, c) is not None:
            v[f] = cell(r, c)

ordered = sorted(vessels.values(), key=lambda v: (0 if v['tx'] > 0 else 1, search_name(v['name']), v['name']))

wb = openpyxl.Workbook()
font = Font(name='Arial', size=10)
bold = Font(name='Arial', size=10, bold=True)
head_font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
blue, purple, grey, green = (PatternFill('solid', fgColor=c) for c in ('1F4E79', '7030A0', '595959', '548235'))
yellow = PatternFill('solid', fgColor='FFF2CC')


def style_header(ws, widths, fills):
    for c, w in enumerate(widths, 1):
        cell_ = ws.cell(row=1, column=c)
        cell_.font = head_font
        cell_.fill = fills[c - 1]
        cell_.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[1].height = 36
    ws.freeze_panes = 'A2'


# ------------------------------------------------------------------------------------------------------------------------------ Petunjuk
ws = wb.active
ws.title = 'Petunjuk'
lines = [
    ('Tujuan', 'Mengisi Vessel Owner, ukuran, GT dan tahun kapal dari Pencarian Data Kapal Ditkapel (https://kapal.dephub.go.id/ditkapel_service/data_kapal/).'),
    ('PENTING', 'Pencarian dilakukan MANUAL. robots.txt situs itu berisi "Disallow: /" untuk semua robot, jadi tidak ada skrip yang mengambil data dari situs.'),
    ('Langkah 1', 'Di sheet Lookup, ambil nama di kolom D ("Cari dengan"). Kerjakan Prioritas 1 dulu (kapal yang dipakai transaksi KLIP).'),
    ('Langkah 2', 'Di situs, isi "Nama Kapal" dengan teks itu, klik Cari Data Kapal.'),
    ('Langkah 3', 'Salin SEMUA baris hasil (satu pencarian bisa mengembalikan beberapa kapal), tempel di sheet Hasil mulai kolom B (kolom No).'),
    ('Langkah 4', 'Isi kolom A di sheet Hasil dengan teks pencarian yang SAMA PERSIS dengan kolom D di Lookup, pada setiap baris yang ditempel.'),
    ('Tidak ada hasil', 'Kalau situs menjawab "Belum ada data", tetap tambah satu baris di Hasil: isi kolom A saja. Itu membedakan "sudah dicari, tidak ada" dari "belum dicari".'),
    ('Jangan', 'Jangan mengubah isi hasil dari situs, dan jangan mengganti nama kolom. Pencocokan dan pemeriksaan dilakukan oleh skrip.'),
    ('Langkah 5', 'Jalankan: python -I docs/scripts/match-vessel-registry.py "<file ini>.xlsx"  -> membuat sheet Registry_Match dan Registry_Summary di file hasil.'),
    ('Dua baris contoh', 'SOLID 10 dan ANAK LAUT 289 di sheet Hasil adalah hasil pencarian sungguhan (13 Okt 2026), bukan data karangan.'),
    ('', ''),
    ('Pemetaan kolom', 'Kolom situs -> kolom baru (bahasa Inggris)'),
    ('Nama Pemilik', 'Vessel Owner (Registry)  - kolom terpisah dari Vessel Owner KLIP: pemilik terdaftar bisa berbeda dari owner komersial (ANAK LAUT 289: KLIP PT. MULIA BORNEO MANDIRI, registry PT. BUANA KENCANA MARITIME)'),
    ('Panjang', 'Vessel Length (Registry)'),
    ('Lebar', 'Vessel Breadth (Registry)'),
    ('Dalam', 'Vessel Depth (Registry)  - DEPTH, bukan draft. Draft (sarat) tidak ada di situs, jadi Vessel Draft tidak ditimpa.'),
    ('LOA', 'Vessel LOA (Registry)'),
    ('GT', 'Vessel GT (Registry)'),
    ('Isi Bersih', 'Vessel NT (Registry)'),
    ('Tahun', 'Vessel Registration Year (Registry)'),
    ('TPK, Call Sign, IMO, Eks Nama Kapal', 'Registration No (TPK), Call Sign, IMO, Former Name  - untuk mencocokkan dan menelusuri'),
]
for r, (a, b) in enumerate(lines, 1):
    ws.cell(row=r, column=1, value=a).font = bold
    c = ws.cell(row=r, column=2, value=b)
    c.font = font
    c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.cell(row=r, column=1).alignment = Alignment(vertical='top', wrap_text=True)
ws.column_dimensions['A'].width = 30
ws.column_dimensions['B'].width = 120

# ------------------------------------------------------------------------------------------------------------------------------ Lookup
ws = wb.create_sheet('Lookup')
heads = ['No', 'Prioritas (1 = dipakai transaksi KLIP)', 'Nama Vessel (KLIP)', 'Cari dengan (ketik ini di situs)', 'Vessel Type', 'Transaksi Klip', 'Pair ID',
         'Vessel Code SAP', 'Vessel Owner (KLIP)', 'Vessel LOA (KLIP)', 'Vessel Draft (KLIP)', 'Vessel Registration Year (KLIP)', 'Status pencarian']
for c, h in enumerate(heads, 1):
    ws.cell(row=1, column=c, value=h)
style_header(ws, [6, 14, 30, 30, 12, 12, 16, 22, 38, 12, 12, 14, 18], [blue] * 4 + [grey] * 8 + [green])
for i, v in enumerate(ordered, 2):
    values = [i - 1, 1 if v['tx'] > 0 else 2, v['name'], search_name(v['name']), v['type'], v['tx'], ', '.join(v['pairs']) or None, v['code'], v['owner'], v['loa'], v['draft'], v['year'],
              f'=IF(COUNTIF(Hasil!$A:$A,D{i})>0,"SUDAH DICARI","BELUM")']
    for c, val in enumerate(values, 1):
        x = ws.cell(row=i, column=c, value=val)
        x.font = font
        if c == 4:
            x.fill = yellow
ws.auto_filter.ref = f'A1:{get_column_letter(len(heads))}{len(ordered) + 1}'

# ------------------------------------------------------------------------------------------------------------------------------ Hasil
ws = wb.create_sheet('Hasil')
ws.cell(row=1, column=1, value='Dicari dengan (sama persis dengan kolom D di Lookup)')
for c, h in enumerate(SITE_COLUMNS, 2):
    ws.cell(row=1, column=c, value=h)
style_header(ws, [34, 6, 30, 22, 14, 14, 36, 22, 10, 10, 10, 10, 8, 10, 12, 8, 14], [green] + [purple] * len(SITE_COLUMNS))
for r, row in enumerate(EXAMPLE_ROWS, 2):
    for c, val in enumerate(row, 1):
        ws.cell(row=r, column=c, value=val if val != '' else None).font = font

wb.move_sheet('Petunjuk', offset=0)
try:
    wb.save(OUT)
except PermissionError:
    OUT = OUT.replace('.xlsx', ' (2).xlsx')
    print('output is open in Excel, saving as', OUT)
    wb.save(OUT)

distinct_terms = len({search_name(v['name']) for v in ordered})
print(f'{len(ordered)} vessels, {distinct_terms} distinct search texts; priority 1: {sum(1 for v in ordered if v["tx"] > 0)}')
print('saved', OUT)
