"""Adds the sheet Vessel_Prefix to a copy of the Vessel Cleanup workbook: the Vessel_Gabungan rows, normalised.

  python -I docs/scripts/build-vessel-prefix.py [source.xlsx] [sheet] [output.xlsx] [--nas "Master Data - CPO 28 Sep 2026.xlsx"] [--combine]
    [--rename "OLD NAME=>NEW NAME"]... [--keep-pair "VESSEL NAME=PAIR ID"]...
    [--latest-pair] [--sap "ZMMLOG_RPT ... .xlsx"]
  (defaults: docs/Vessel Cleanup (Jovin, Klip, SAP) v6.xlsx / Vessel_Gabungan -> ... v7.xlsx, v8.xlsx with --nas, v9.xlsx with --combine; when the output is open in Excel it is saved as '(2)')

--nas carries what people filled in on the NAS sheet 'Vessel - Pair' (Vessel Owner, Vessel Type, Charter Type, Vessel Capacity) into the list BEFORE the name rules run, so a
Vessel Type given in the NAS also settles a missing prefix. The NAS wins: where it has a value it replaces ours; where it is blank ours stays. Its rows are realigned first
(vessel_nas_merge.py); the vessels it has that this list has not are listed in the sheet NAS_Only, not added.

--rename fixes a name by decision (a typo) before anything else, so the corrected name then merges like any other. --keep-pair decides that a vessel which sits in several pairs is
ONE vessel belonging to the given Pair ID: its other pair rows are merged into that row (needs --combine).

--latest-pair (with --combine) turns every vessel that still sits in several pair rows into ONE row that keeps the Pair ID of its latest contract (vessel_pair_dates.py: the workbook's
Vessel_Pair_Link, and the SAP export given with --sap matched by the raw SAP name of each pair). Without a date the lowest Pair ID is kept and the row says so.

--combine merges the rows that are one vessel (vessel_combine.py): rows without a Pair ID fold into the row(s) of the same vessel that have one, or into each other, so the same
name is not listed twice. What was merged is in the sheet Combine_Log; names that look alike but were NOT merged are in Near_Duplicates.

Columns A:Q keep the layout of the NAS sheet "Vessel - Klip (Final Validation)" so they can be pasted over it; R:V are info only.
Rules: every value in A:Q upper-case. Vessel name: the prefix is added from Vessel Type (BARGE -> BG., TUG BOAT -> TB., TANKER -> MT., SPOB -> SPOB.),
a prefix without its dot or glued to the name is rewritten as 'XX. NAME', TK. becomes BG.; KLM. / KM. / MV. only get their space and are flagged.
Vessel Owner: 'PT.' moved to the front, 'PT.' added when missing, 'PT' / 'P.T.' written as 'PT.'; CV. is kept (moved to the front); two companies joined by ' / '
are treated one by one; a value cut at SAP's 35 characters ending ' P' is restored to PT.; an incoterm (CIF) in the owner column is NOT a company: left and flagged.
Rows that need a decision are yellow, changed rows green. Nothing is dropped: every Pair ID and every other column is copied as it was.
Needs openpyxl.
"""
import difflib
import os
import re
import sys
import collections
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_args = sys.argv[1:]
COMBINE = '--combine' in _args
if COMBINE:
    _args.remove('--combine')
LATEST_PAIR = '--latest-pair' in _args
if LATEST_PAIR:
    _args.remove('--latest-pair')


def _take_all(flag):
    values = []
    while flag in _args:
        i = _args.index(flag)
        values.append(_args[i + 1])
        del _args[i:i + 2]
    return values


SAP_FILE = (_take_all('--sap') or [None])[0]
RENAMES = dict(v.split('=>', 1) for v in _take_all('--rename'))
KEEP_PAIRS = [tuple(v.rsplit('=', 1)) for v in _take_all('--keep-pair')]
NAS = None
if '--nas' in _args:
    _i = _args.index('--nas')
    NAS = _args[_i + 1]
    del _args[_i:_i + 2]
SRC = _args[0] if len(_args) > 0 else 'D:/Project/Klip/docs/Vessel Cleanup (Jovin, Klip, SAP) v6.xlsx'
SHEET_IN = _args[1] if len(_args) > 1 else 'Vessel_Gabungan'
OUT = _args[2] if len(_args) > 2 else f"D:/Project/Klip/docs/Vessel Cleanup (Jovin, Klip, SAP) {'v9' if COMBINE else 'v8' if NAS else 'v7'}.xlsx"
SHEET_OUT = 'Vessel_Prefix'

# the owner / name rules live in vessel_common.py (shared with the registry scripts); -I drops the script directory from sys.path, so add it back
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from vessel_common import fix_vessel_name, key, norm_owner, parse_number, strip_prefix  # noqa: E402
from vessel_nas_merge import load_nas_edits, pick  # noqa: E402
from vessel_combine import combine, keep_latest_pairs, keep_one_pair, near_pairs, split_name  # noqa: E402
from vessel_pair_dates import load_pair_dates  # noqa: E402


wb = openpyxl.load_workbook(SRC)
ws_in = wb[SHEET_IN]
data = [list(r) for r in ws_in.iter_rows(values_only=True)]
header, rows = data[0], data[1:]
ncol = min(len(header), 17)  # A:Q (A:O NAS, P Vessel Code SAP, Q Sumber) = the layout of 'Vessel - Klip (Final Validation)'

NAS_EDITS, NAS_STATS = ({}, collections.Counter())
NAS_BASE = collections.defaultdict(list)
if NAS:
    NAS_EDITS, NAS_STATS = load_nas_edits(NAS)
    for k, e in NAS_EDITS.items():
        NAS_BASE[key(strip_prefix(e['name']))].append(k)
_PREFIX_TOKEN = re.compile(r'^(BG|TB|MT|SPOB)\.')

# (field, label, index in this list, how to read ours)
NAS_FIELDS = (
    ('type', 'Vessel Type', 7, lambda v: str(v).strip().upper() if v not in (None, '') else None),
    ('owner', 'Vessel Owner', 5, lambda v: norm_owner(v)[0]),
    ('charter', 'Charter Type', 10, lambda v: str(v).strip().upper() if v not in (None, '') else None),
    ('capacity', 'Vessel Capacity', 6, lambda v: parse_number(v)),
)
nas_note_by_row = {}
nas_matched = set()
nas_changes = collections.Counter()


def find_nas_entry(name, vtype):
    canon, _ = fix_vessel_name(name, vtype)
    k = key(canon)
    if k in NAS_EDITS:
        return k
    cands = NAS_BASE.get(key(strip_prefix(name)), [])
    if len(cands) != 1:
        return None
    # the same base name under two different prefixes is two vessels (a tug and a barge), never one
    a, b = _PREFIX_TOKEN.match(canon.upper()), _PREFIX_TOKEN.match(NAS_EDITS[cands[0]]['name'].upper())
    return cands[0] if not (a and b and a.group(1) != b.group(1)) else None


rename_note_by_row = {}
if RENAMES:
    by_key = {key(a): b.strip() for a, b in RENAMES.items()}
    for r in rows:
        if r[1] is not None and key(r[1]) in by_key:
            rename_note_by_row[id(r)] = f"Nama diganti sesuai keputusan: {str(r[1]).strip()} -> {by_key[key(r[1])]}"
            r[1] = by_key[key(r[1])]

# pass 1: which NAS entry (if any) each row belongs to
match_of, rename_of, nas_choice = {}, {}, {}
if NAS:
    for r in rows:
        nm = str(r[1]).strip() if r[1] is not None else ''
        match_of[id(r)] = find_nas_entry(nm, r[7]) if nm else None
    nas_matched.update(v for v in match_of.values() if v)
    # pass 2: a name the NAS spells differently (a typo fixed there): the closest name of this list that the NAS did not claim, when it is clearly the same one
    free = {}
    for r in rows:
        if r[1] and match_of[id(r)] is None:
            free.setdefault(key(fix_vessel_name(str(r[1]).strip(), r[7])[0]), []).append(r)
    for nk in [k for k in NAS_EDITS if k not in nas_matched]:
        scored = sorted(((difflib.SequenceMatcher(None, nk, fk).ratio(), fk) for fk in free), reverse=True)
        if scored and scored[0][0] >= 0.9 and (len(scored) == 1 or scored[1][0] < 0.9):
            for r in free[scored[0][1]]:
                match_of[id(r)] = nk
                rename_of[id(r)] = NAS_EDITS[nk]['name']
            nas_matched.add(nk)

out = []
for r in rows:
    name = str(r[1]).strip() if r[1] is not None else ''
    nk = match_of.get(id(r))
    if nk is not None:
        e = NAS_EDITS[nk]
        parts = []
        if id(r) in rename_of:
            parts.append(f'Nama: {name} -> {rename_of[id(r)]} (ejaan di NAS)')
            nas_changes[('nama', 'diganti')] += 1
            name = rename_of[id(r)]
            r[1] = name
        for field, label, idx, read in NAS_FIELDS:
            ours = read(r[idx])
            if (nk, field) not in nas_choice:  # one value per vessel, decided once, so every row of the vessel agrees
                nas_choice[(nk, field)] = pick(e[field], ours)
            value, conflict = nas_choice[(nk, field)]
            if isinstance(value, float) and value.is_integer():
                value = int(value)
            if value is not None and value != ours and not (isinstance(ours, float) and ours == value):
                parts.append(f"{label}: {ours if ours is not None else '(kosong)'} -> {value}" + (f' [{conflict}]' if conflict else ''))
                nas_changes[(field, 'diisi' if ours is None else 'diganti')] += 1
                r[idx] = value
        if parts:
            nas_note_by_row[id(r)] = '; '.join(parts)
    new, note = fix_vessel_name(name, r[7])
    out.append((r, name, new, note))

combine_note_by_row, combine_log, combine_near = {}, [], []
rows_before = len(out)
if COMBINE:
    items = [{'r': r, 'old': name, 'new': new} for r, name, new, note in out]
    kept, combine_log, combine_near = combine(items)
    for vessel, pair_id in KEEP_PAIRS:
        kept, entry = keep_one_pair(kept, vessel, pair_id, 'keputusan: satu vessel, pair dipilih manual')
        if entry:
            combine_log.append(entry)
    if LATEST_PAIR:
        pair_dates, pair_source = load_pair_dates(wb, SAP_FILE)
        kept, latest_log = keep_latest_pairs(kept, pair_dates, pair_source, skip={key(v) for v, _ in KEEP_PAIRS})
        combine_log.extend(latest_log)
        print(f'latest-pair: {len(latest_log)} vessels reduced to one row')
    out = []
    for it in kept:
        r = it['r']
        new, note = fix_vessel_name(str(r[1]).strip() if r[1] is not None else '', r[7])
        if it.get('combined'):
            combine_note_by_row[id(r)] = it['combined']
        out.append((r, it['old'], new, note))

# the same vessel under two spellings: after the prefix, two rows carry one name
count = collections.Counter(key(new) for _, _, new, _ in out if new)
first_pair = {}
for r, _, new, _ in out:
    if new and r[0]:
        first_pair.setdefault(key(new), (r[0], new))

if SHEET_OUT in wb.sheetnames:
    del wb[SHEET_OUT]
ws = wb.create_sheet(SHEET_OUT, 1)
font = Font(name='Arial', size=10)
head_font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
blue, purple, grey = (PatternFill('solid', fgColor=c) for c in ('1F4E79', '7030A0', '595959'))
changed_fill = PatternFill('solid', fgColor='E2EFDA')
decide_fill = PatternFill('solid', fgColor='FFF2CC')
dupe_fill = PatternFill('solid', fgColor='FCE4D6')

heads = list(header[:ncol]) + ['Nama sebelum (info)', 'Prefix: apa yang dilakukan (info)', 'Nama sama dengan baris lain (info)', 'Owner sebelum (info)', 'Owner: apa yang dilakukan (info)']
if NAS:
    heads.append('NAS: apa yang diubah (info)')
if COMBINE:
    heads.append('Digabung (info)')
for c, h in enumerate(heads, 1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.font = head_font
    cell.fill = blue if c <= 4 else (purple if c <= 16 else grey)
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
ws.row_dimensions[1].height = 42

stats = collections.Counter()
for i, (r, old, new, note) in enumerate(out, 2):
    values = list(r[:ncol])
    values[1] = new or None
    owner_new, owner_note = norm_owner(r[5])
    values[5] = owner_new
    values = [v.upper() if isinstance(v, str) else v for v in values]
    dup = ''
    if new and count[key(new)] > 1:
        group = [o for o in out if o[2] and key(o[2]) == key(new)]
        loose = [o for o in group if not o[0][0]]  # rows with no Pair ID
        # a vessel that sits in several pairs is normal; the case to look at is a row WITHOUT a Pair ID that is the same vessel as another row
        if loose:
            pair_ids = sorted({o[0][0] for o in group if o[0][0]})
            if not r[0]:
                dup = 'Sama dengan baris ber-Pair ID: ' + ', '.join(pair_ids) if pair_ids else f'{len(group)} baris tanpa Pair ID bernama sama'
            else:
                dup = 'Ada baris tanpa Pair ID dengan nama sama'
    row_fill = None
    if new != old and note.startswith(('Prefix', 'Titik', 'TK.')) and 'tidak cocok' not in note:
        row_fill = changed_fill
        stats['diubah'] += 1
    if note.startswith('Perlu diputuskan') or 'belum ada aturannya' in note or 'tidak cocok' in note or note == 'Nama kosong':
        row_fill = decide_fill
        stats['perlu diputuskan'] += 1
    if dup:
        stats['nama kembar'] += 1
    for c, v in enumerate(values, 1):
        cell = ws.cell(row=i, column=c, value=v)
        cell.font = font
        if row_fill:
            cell.fill = row_fill
    ws.cell(row=i, column=ncol + 1, value=old if old != new else None).font = font
    ws.cell(row=i, column=ncol + 2, value=note or None).font = font
    ws.cell(row=i, column=ncol + 4, value=(str(r[5]).strip() if r[5] is not None and owner_new != str(r[5]).strip() else None)).font = font
    ws.cell(row=i, column=ncol + 5, value=owner_note or None).font = font
    if owner_note:
        stats['owner diubah'] += 1
    if owner_note.startswith('Bukan nama'):
        stats['owner bukan perusahaan'] += 1
    if NAS:
        nas_cell = ws.cell(row=i, column=ncol + 6, value=nas_note_by_row.get(id(r)))
        nas_cell.font = font
        if nas_note_by_row.get(id(r)):
            nas_cell.fill = changed_fill
            stats['baris diubah NAS'] += 1
    if COMBINE:
        cm = ws.cell(row=i, column=ncol + 6 + (1 if NAS else 0), value=combine_note_by_row.get(id(r)))
        cm.font = font
        if combine_note_by_row.get(id(r)):
            cm.fill = changed_fill
    d = ws.cell(row=i, column=ncol + 3, value=dup or None)
    d.font = font
    if dup:
        d.fill = dupe_fill

widths = [10, 30, 12, 12, 14, 38, 12, 12, 9, 12, 11, 11, 11, 16, 14, 24, 52, 30, 54, 44, 40, 60, 90, 90]
for c, w in enumerate(widths[:len(heads)], 1):
    ws.column_dimensions[get_column_letter(c)].width = w
ws.freeze_panes = 'C2'
ws.auto_filter.ref = f'A1:{get_column_letter(len(heads))}{len(out) + 1}'
if NAS:
    if 'NAS_Only' in wb.sheetnames:
        del wb['NAS_Only']
    wo = wb.create_sheet('NAS_Only', 2)
    only_heads = ['Nama Vessel (NAS)', 'Baris di NAS', 'Vessel Owner', 'Vessel Type', 'Charter Type', 'Vessel Capacity']
    for c, h in enumerate(only_heads, 1):
        x = wo.cell(row=1, column=c, value=h)
        x.font, x.fill = head_font, grey
    n_only = 0
    for k, e in NAS_EDITS.items():
        if k in nas_matched:
            continue
        n_only += 1
        vals = [e['name'], e['rows']] + [' | '.join(str(v) for v in e[f]) or None for f in ('owner', 'type', 'charter', 'capacity')]
        for c, v in enumerate(vals, 1):
            wo.cell(row=n_only + 1, column=c, value=v.upper() if isinstance(v, str) else v).font = font
    for c, w in enumerate([34, 12, 44, 14, 14, 16], 1):
        wo.column_dimensions[get_column_letter(c)].width = w
if COMBINE:
    for name in ('Combine_Log', 'Near_Duplicates'):
        if name in wb.sheetnames:
            del wb[name]
    wl = wb.create_sheet('Combine_Log', 2)
    for c, h in enumerate(['Nama Vessel (dipertahankan)', 'Pair ID', 'Baris yang digabung (nama asal)', 'Diisi dari baris lain', 'Beda nilai (nilai baris yang dipertahankan dipakai)'], 1):
        x = wl.cell(row=1, column=c, value=h)
        x.font, x.fill = head_font, grey
    for i, e in enumerate(combine_log, 2):
        for c, v in enumerate([e['name'], ', '.join(p for p in e['pairs'] if p) or None, ' | '.join(e['merged']), ', '.join(e['filled']) or None, ' | '.join(dict.fromkeys(e['conflicts'])) or None], 1):
            wl.cell(row=i, column=c, value=v.upper() if isinstance(v, str) else v).font = font
    for c, w in enumerate([34, 22, 50, 50, 70], 1):
        wl.column_dimensions[get_column_letter(c)].width = w

    # names that look alike but were NOT merged: a spelling variant (SAMUDRA / SAMUDERA) or a different vessel; the same numbers and Roman numerals are required
    wn = wb.create_sheet('Near_Duplicates', 3)
    for c, h in enumerate(['Nama Vessel A', 'Pair ID A', 'Nama Vessel B', 'Pair ID B', 'Kemiripan', 'Catatan'], 1):
        x = wn.cell(row=1, column=c, value=h)
        x.font, x.fill = head_font, grey
    firsts = {}
    for r, _, new, _ in out:
        firsts.setdefault(new, r[0])
    near_rows = [(na, pa, nb, pb, ratio, 'ejaan mirip, belum digabung') for na, pa, nb, pb, ratio in near_pairs(list(firsts.items()))]
    for a_name, why in combine_near:
        near_rows.append((a_name, None, ' / '.join(why), None, None, 'tanpa prefix, cocok dengan lebih dari satu vessel berprefix (BG dan TB)'))
    for i, row in enumerate(sorted(near_rows, key=lambda x: -(x[4] or 0)), 2):
        for c, v in enumerate(row, 1):
            wn.cell(row=i, column=c, value=v.upper() if isinstance(v, str) else v).font = font
    for c, w in enumerate([34, 12, 34, 12, 11, 60], 1):
        wn.column_dimensions[get_column_letter(c)].width = w
try:
    wb.save(OUT)
except PermissionError:
    OUT = OUT.replace('.xlsx', ' (2).xlsx')
    print('output is open in Excel, saving as', OUT)
    wb.save(OUT)

by_note = collections.Counter()
for _, _, _, note in out:
    k = note.split(' dari ')[0].split(' - ')[0] if note else '(sudah benar)'
    by_note[re.sub(r'Prefix (BG|TB|MT|SPOB)\. ditambah', 'Prefix ditambah', k)] += 1
print('rows', len(out))
for k, v in by_note.most_common():
    print(f'  {v:4}  {k}')
print('changed:', stats['diubah'], '| need a decision:', stats['perlu diputuskan'], '| rows sharing a name with another row:', stats['nama kembar'])
print('owner rows changed:', stats['owner diubah'], '| not a company:', stats['owner bukan perusahaan'])
if NAS:
    print('NAS rows read:', dict(NAS_STATS))
    print('NAS vessels:', len(NAS_EDITS), '| matched to this list:', len(nas_matched), '| only in the NAS (sheet NAS_Only):', len(NAS_EDITS) - len(nas_matched))
    print('rows of this list changed from the NAS:', stats['baris diubah NAS'], '| cell changes:', {f'{a} {b}': n for (a, b), n in sorted(nas_changes.items())})
if COMBINE:
    print(f'combine: {rows_before} -> {len(out)} rows ({rows_before - len(out)} merged away) | merge groups: {len(combine_log)} | near duplicates listed: {len(near_rows)}')
print('saved', OUT)
