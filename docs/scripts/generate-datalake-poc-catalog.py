#!/usr/bin/env python3
"""Build KLIP Data Lake POC Excel catalog from live Postgres schema dumps (CSV)."""

from __future__ import annotations

import csv
import re
from collections import defaultdict
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

ROOT = Path(__file__).resolve().parents[2]
CSV_DIR = ROOT / "docs"
OUT = ROOT / "docs" / "datalake" / "KLIP-Data-Lake-POC-Catalog.xlsx"
AS_OF = date.today().isoformat()

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True, name="Calibri", size=10)
CELL_FONT = Font(name="Calibri", size=10)
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)
P0 = PatternFill("solid", fgColor="C6EFCE")
P1 = PatternFill("solid", fgColor="FFEB9C")
P2 = PatternFill("solid", fgColor="DDEBF7")
EX = PatternFill("solid", fgColor="F8CBAD")
SENS = PatternFill("solid", fgColor="F4B183")

# domain, grain, lake_layer, ingest_priority, description
TABLE_META: dict[str, tuple[str, str, str, str, str]] = {
    "sap_data_imports": (
        "SAP ingest",
        "1 row per Excel/file import session",
        "BRONZE — source of ingest sessions",
        "P0",
        "SAP MASTER v2 upload session (manual or scheduler). Status, row counts, file hash, trust flag.",
    ),
    "sap_raw_data": (
        "SAP ingest",
        "1 row per Excel row in an import",
        "BRONZE — raw payload",
        "P0",
        "Row-level raw JSON from SAP Excel plus parse errors. Highest volume; keep for replay/lineage.",
    ),
    "sap_processed_data": (
        "SAP ingest",
        "1 row per tri-key (contract + PO + STO) per processed import lineage",
        "SILVER — normalized SAP",
        "P0",
        "Normalized SAP row. Identity columns plus JSONB data (contract/shipment/trucking/vessel/payment/quality/raw). Primary commercial source after ingest.",
    ),
    "sap_import_failures": (
        "SAP ingest",
        "1 row per failed Excel row",
        "BRONZE — data quality",
        "P1",
        "Failed SAP rows and error messages. Use for ingest quality KPIs.",
    ),
    "sap_auto_import_files": (
        "SAP ingest",
        "1 row per scheduled file SHA-256",
        "APPLICATION_OPS",
        "P2",
        "Scheduler skip-list so identical Original-folder files are not re-imported.",
    ),
    "sap_field_mappings": (
        "SAP ingest",
        "mapping row",
        "REFERENCE",
        "P2",
        "Optional field-mapping configuration (often unused; MASTER v2 mapping is in code).",
    ),
    "sap_presence_audit": (
        "SAP ingest",
        "audit event when a tri-key disappears/reappears",
        "SILVER — lineage",
        "P1",
        "Tracks SAP presence vs WITHDRAWN after trusted imports.",
    ),
    "contracts": (
        "Commercial",
        "1 row per PO (primary commercial grain); STO list is in contract_stos",
        "GOLD — contract master",
        "P0",
        "Contract/PO master distributed from SAP. Buyer, supplier, product, qty, dates, plant, presence. Outstanding qty is NOT stored here; it is computed from SAP/snapshots.",
    ),
    "contract_stos": (
        "Commercial",
        "1 row per STO on a contract/PO",
        "GOLD — STO list",
        "P0",
        "STO numbers linked to a contract. Required for multi-STO POs (search, trucking, SEA matching).",
    ),
    "contract_latest_spd_snapshot": (
        "Commercial snapshot",
        "1 row per latest SAP processed identity for list UI",
        "GOLD derived — rebuildable",
        "P1",
        "Denormalized latest sap_processed_data fields for Contracts list performance. Do not treat as independent source; refresh from SPD.",
    ),
    "contract_latest_spd_snapshot_meta": (
        "Commercial snapshot",
        "singleton refresh metadata",
        "APPLICATION_OPS",
        "P2",
        "Last refresh time/status for contract_latest_spd_snapshot.",
    ),
    "contract_performance_snapshot": (
        "Commercial snapshot",
        "1 row per contract for Contract Performance page",
        "GOLD derived — rebuildable",
        "P1",
        "Precomputed Contract Performance KPIs (delivered/OS, GR status, SAP presence).",
    ),
    "contract_performance_snapshot_meta": (
        "Commercial snapshot",
        "singleton refresh metadata",
        "APPLICATION_OPS",
        "P2",
        "Refresh metadata for contract_performance_snapshot.",
    ),
    "contract_qty_move_snapshot": (
        "Commercial snapshot",
        "1 row per contract qty-move overlay",
        "GOLD derived — rebuildable",
        "P1",
        "Delivery/receive qty overlay used by Contracts (Open→KLIP / Close→SAP, B2B child SUM rules).",
    ),
    "contract_qty_move_snapshot_meta": (
        "Commercial snapshot",
        "singleton refresh metadata",
        "APPLICATION_OPS",
        "P2",
        "Refresh metadata for contract_qty_move_snapshot.",
    ),
    "contract_sto_agg_snapshot": (
        "Commercial snapshot",
        "1 row per contract with aggregated STO numbers/qty",
        "GOLD derived — rebuildable",
        "P1",
        "STO aggregation for contracts list search/filters. Missing STO identity here means OS/delivered can look wrong.",
    ),
    "contract_sto_agg_snapshot_meta": (
        "Commercial snapshot",
        "singleton refresh metadata",
        "APPLICATION_OPS",
        "P2",
        "Refresh metadata for contract_sto_agg_snapshot.",
    ),
    "b2b_ending_child_snapshot": (
        "Commercial snapshot",
        "1 row per B2B origin PO overlay",
        "GOLD derived — rebuildable",
        "P1",
        "Child-PO overlay for B2B origin (empty Contract Reff PO): buyer/plant/GR STO/discharge from children.",
    ),
    "b2b_ending_child_snapshot_meta": (
        "Commercial snapshot",
        "singleton refresh metadata",
        "APPLICATION_OPS",
        "P2",
        "Refresh metadata for b2b_ending_child_snapshot.",
    ),
    "shipments": (
        "SEA execution",
        "1 row per shipment (usually 1 active SEA op per contract/STO)",
        "GOLD — SEA shipment",
        "P0",
        "Sea shipment execution: vessel, ports, quantities (kg), milestone dates (ETA/ATA), status. Port-level detail is in vessel_loading_ports.",
    ),
    "vessel_loading_ports": (
        "SEA execution",
        "1 row per loading or discharge port on a shipment",
        "GOLD — port timeline",
        "P0",
        "Per-port SEA timeline and quantities (multi-loading ports + discharge). Authoritative for port-level ETA/ATA vs shipment-level summary.",
    ),
    "shipment_ata_overrides": (
        "SEA execution",
        "1 row per manual ATA override",
        "GOLD — exception overlay",
        "P1",
        "User/KLIP ATA overrides that must win over SAP-derived dates.",
    ),
    "pre_planned_groups": (
        "SEA planning",
        "1 vessel/parcel planning group",
        "GOLD — planning",
        "P1",
        "Preplanned vessel groups (Unplanned→Preplanned workflow).",
    ),
    "pre_planned_group_members": (
        "SEA planning",
        "1 contract/shipment membership in a group",
        "GOLD — planning",
        "P1",
        "Members of a pre-planned group.",
    ),
    "pre_planned_parcel_capacity": (
        "SEA planning",
        "capacity config",
        "REFERENCE",
        "P2",
        "Parcel/vessel capacity helper for preplanning.",
    ),
    "pre_planned_rebuild_log": (
        "SEA planning",
        "rebuild event",
        "APPLICATION_OPS",
        "P2",
        "Log of preplanned-group rebuild jobs.",
    ),
    "shipment_list_stage_snapshot": (
        "SEA snapshot",
        "1 row per shipment list card/stage",
        "GOLD derived — rebuildable",
        "P1",
        "Shipments page list/stage snapshot for UI performance.",
    ),
    "shipment_pipeline_daily_summary": (
        "SEA snapshot",
        "1 row per day × pipeline stage",
        "GOLD derived — rebuildable",
        "P1",
        "Daily pipeline counts/qty for Shipments status cards.",
    ),
    "shipment_pipeline_vessel_stage_daily": (
        "SEA snapshot",
        "1 row per day × vessel × stage",
        "GOLD derived — rebuildable",
        "P2",
        "Vessel-level daily pipeline stage facts.",
    ),
    "loading_ports": (
        "SEA execution",
        "legacy multi-port row",
        "LEGACY",
        "P2",
        "Older multi-port table. Prefer vessel_loading_ports for current SEA timelines.",
    ),
    "trucking_operations": (
        "LAND execution",
        "1 row per trucking operation (typically 1 active op per contract)",
        "GOLD — trucking header",
        "P0",
        "Land/trucking operation header: status, OA, dates, quantities, plant/region, KLIP-edited provenance.",
    ),
    "trucking_daily_actuals": (
        "LAND execution",
        "1 row per weighbridge/daily actual (often per STO/date)",
        "GOLD — weighbridge facts",
        "P0",
        "Daily actuals / WB (weighbridge) quantities and dates. Grain can include blank STO; used for delivery/receive.",
    ),
    "trucking_realizations": (
        "LAND execution",
        "1 row per realization leg/line",
        "GOLD — trucking legs",
        "P1",
        "Supporting trucking legs/realizations under an operation.",
    ),
    "trucking_wb_imports": (
        "LAND ingest",
        "1 row per WB Excel import",
        "BRONZE — WB ingest",
        "P1",
        "Weighbridge Excel import sessions.",
    ),
    "trucking_list_stage_snapshot": (
        "LAND snapshot",
        "1 row per trucking list card",
        "GOLD derived — rebuildable",
        "P1",
        "Trucking page list snapshot (filters, OS, late indicator). Rebuild from operations + SAP + WB.",
    ),
    "trucking_pipeline_daily_summary": (
        "LAND snapshot",
        "1 row per day × trucking stage",
        "GOLD derived — rebuildable",
        "P1",
        "Daily trucking pipeline summary for status cards.",
    ),
    "trucking_stage_snapshot_delta_log": (
        "LAND snapshot",
        "delta event",
        "APPLICATION_OPS",
        "P2",
        "Change log for trucking stage snapshot refreshes.",
    ),
    "payments": (
        "Finance",
        "1 row per payment/invoice line on a contract",
        "GOLD — payments",
        "P0",
        "Invoice/payment due vs paid, dates, amounts. Aging is computed (due vs payoff).",
    ),
    "mv_contract_payment_dates": (
        "Finance",
        "1 row per contract payment-date rollup",
        "GOLD derived — materialized view",
        "P1",
        "Materialized view of payment dates per contract for finance dashboards. Refresh with the app job.",
    ),
    "settlement_invoice_summaries": (
        "Finance",
        "settlement summary row",
        "GOLD",
        "P2",
        "Settlement invoice summary (may be sparsely used).",
    ),
    "quality_surveys": (
        "Quality",
        "1 row per survey (loading/discharge) on a shipment",
        "GOLD — quality",
        "P1",
        "COA/survey results (FFA, moisture, etc.). High volume locally because SAP quality arrays are exploded.",
    ),
    "surveyors": (
        "Quality",
        "1 row per surveyor master",
        "REFERENCE",
        "P2",
        "Surveyor master (charges/names).",
    ),
    "claim_mutu_imports": (
        "Claims",
        "1 import session",
        "BRONZE",
        "P1",
        "Claim Mutu (quality claim) Excel import header.",
    ),
    "claim_mutu_rows": (
        "Claims",
        "1 Excel claim row",
        "GOLD overlay",
        "P1",
        "Claim Mutu overlay rows (not a replacement for SAP quality).",
    ),
    "claim_susut_imports": (
        "Claims",
        "1 import session",
        "BRONZE",
        "P1",
        "Claim Susut (loss/shortage) Excel import header.",
    ),
    "claim_susut_rows": (
        "Claims",
        "1 Excel claim row",
        "GOLD overlay",
        "P1",
        "Claim Susut overlay; joined to trucking/SAP in the UI, not a full SPD scan.",
    ),
    "documents": (
        "Documents",
        "1 uploaded file",
        "GOLD — documents metadata",
        "P1",
        "Uploaded files linked to contract/shipment/payment/trucking. File bytes are on volume/NAS, not in this table.",
    ),
    "commercial_document_files": (
        "Documents",
        "1 commercial doc file",
        "GOLD",
        "P1",
        "Commercial documents module files.",
    ),
    "commercial_document_history": (
        "Documents",
        "history event",
        "GOLD",
        "P2",
        "Commercial document status/history.",
    ),
    "remarks": (
        "Collaboration",
        "1 remark on an entity",
        "GOLD",
        "P2",
        "User remarks on contracts/shipments/etc.",
    ),
    "alerts": (
        "Collaboration",
        "1 alert",
        "GOLD",
        "P2",
        "Operational alerts.",
    ),
    "audit_logs": (
        "Control",
        "1 user/system change event",
        "GOLD — audit",
        "P1",
        "Who changed what/when. Needed for lake lineage and SoD, not for logistics KPIs.",
    ),
    "user_activity_events": (
        "Control",
        "1 UI activity event",
        "APPLICATION_OPS",
        "P2",
        "In-app activity stream (page-level). High churn.",
    ),
    "users": (
        "IAM",
        "1 application user",
        "RESTRICTED",
        "P1",
        "KLIP users. Contains password hashes and PII — mask in the lake.",
    ),
    "roles": (
        "IAM",
        "1 role",
        "REFERENCE",
        "P1",
        "RBAC roles (ADMIN, TRADING, LOGISTICS, FINANCE, MANAGEMENT, SUPPORT, …).",
    ),
    "permissions": (
        "IAM",
        "1 permission key",
        "REFERENCE",
        "P1",
        "Permission catalog (page.* and action keys).",
    ),
    "role_permissions": (
        "IAM",
        "role × permission",
        "REFERENCE",
        "P1",
        "Grants: view/create/edit/delete per role.",
    ),
    "user_plants": (
        "IAM",
        "user × plant/region scope",
        "REFERENCE",
        "P1",
        "Data-scope filter: which plants/regions a user may see.",
    ),
    "user_products": (
        "IAM",
        "user × product scope",
        "REFERENCE",
        "P2",
        "Optional product-level data scope.",
    ),
    "user_preferences": (
        "IAM",
        "1 user preference blob/row",
        "APPLICATION_OPS",
        "P2",
        "UI preferences (columns, filters). Low lake value.",
    ),
    "user_sto_contract_assignments": (
        "IAM",
        "assignment row",
        "GOLD",
        "P2",
        "STO/contract assignments to users (workflow).",
    ),
    "user_data_inputs": (
        "IAM",
        "manual data input",
        "GOLD",
        "P2",
        "Generic user-entered overlay store.",
    ),
    "master_plants": (
        "Master data",
        "1 plant / group-plant",
        "GOLD — reference",
        "P0",
        "Plant and region/group mapping (e.g. EU23 / TANJUNG PURA). Used for filters and user scope.",
    ),
    "master_vessels": (
        "Master data",
        "1 canonical vessel",
        "GOLD — reference",
        "P0",
        "Canonical vessel identity (name/code) including Jovin import fields.",
    ),
    "master_vessel_code_aliases": (
        "Master data",
        "1 alias → canonical vessel",
        "GOLD — reference",
        "P1",
        "Vessel code aliases for matching SAP/Jovin codes to master_vessels.",
    ),
    "master_loading_ports": (
        "Master data",
        "1 port",
        "GOLD — reference",
        "P1",
        "Loading port master.",
    ),
    "vessel_master": (
        "Master data",
        "legacy vessel master row",
        "LEGACY",
        "P2",
        "Older vessel master. Prefer master_vessels.",
    ),
    "suppliers": (
        "Master data",
        "1 supplier",
        "GOLD — reference",
        "P1",
        "Supplier master (often sourced/synced from SAP names).",
    ),
    "supplier_groups": (
        "Master data",
        "1 vendor group",
        "REFERENCE",
        "P2",
        "Vendor/supplier grouping.",
    ),
    "products": (
        "Master data",
        "1 product",
        "GOLD — reference",
        "P1",
        "Product configuration used by filters/AI.",
    ),
    "eta_patterns": (
        "Master data",
        "ETA pattern",
        "REFERENCE",
        "P2",
        "ETA estimation patterns (AI/rules).",
    ),
    "vessel_patterns": (
        "Master data",
        "vessel pattern",
        "REFERENCE",
        "P2",
        "Vessel behaviour patterns for planning/AI.",
    ),
    "ai_insights": (
        "AI",
        "1 generated insight",
        "APPLICATION_OPS",
        "P2",
        "Cached Gemini/AI insights per filter set.",
    ),
    "dashboard_ai_insights": (
        "AI",
        "dashboard insight cache",
        "APPLICATION_OPS",
        "P2",
        "Dashboard-level AI insight cache.",
    ),
    "agent_ai_memory": (
        "AI",
        "1 Q/A memory",
        "RESTRICTED",
        "P2",
        "KLIP Agent AI Q/A memory. May contain operational text; not a logistics fact table.",
    ),
    "agent_ai_lessons": (
        "AI",
        "1 team-wide lesson",
        "APPLICATION_OPS",
        "P2",
        "Curated lessons the Agent AI may apply team-wide.",
    ),
    "ai_klip_agent_activity_logs": (
        "AI",
        "1 agent activity",
        "APPLICATION_OPS",
        "P2",
        "Agent AI activity log.",
    ),
    "data_validation_rules": (
        "Control",
        "1 validation rule",
        "REFERENCE",
        "P2",
        "Configurable validation rules.",
    ),
    "pipeline_summary_refresh_meta": (
        "Ops",
        "pipeline job metadata",
        "APPLICATION_OPS",
        "P2",
        "Refresh timestamps for shipment/trucking pipeline summaries.",
    ),
    "schema_migrations": (
        "Ops",
        "1 applied migration",
        "EXCLUDE from analytics",
        "Exclude",
        "App schema version log. Useful for lake schema drift checks, not for KPIs.",
    ),
    "cleanup_audit_062": (
        "Ops cleanup",
        "one-off cleanup residue",
        "EXCLUDE from analytics",
        "Exclude",
        "Historical data-fix audit. Do not ingest as a domain table.",
    ),
    "cleanup_audit_077": (
        "Ops cleanup",
        "one-off cleanup residue",
        "EXCLUDE from analytics",
        "Exclude",
        "Historical data-fix audit. Do not ingest as a domain table.",
    ),
    "cleanup_audit_078": (
        "Ops cleanup",
        "one-off cleanup residue",
        "EXCLUDE from analytics",
        "Exclude",
        "Historical data-fix audit. Do not ingest as a domain table.",
    ),
    "cleanup_audit_115": (
        "Ops cleanup",
        "one-off cleanup residue",
        "EXCLUDE from analytics",
        "Exclude",
        "Historical data-fix audit. Do not ingest as a domain table.",
    ),
    "cleanup_audit_116": (
        "Ops cleanup",
        "one-off cleanup residue",
        "EXCLUDE from analytics",
        "Exclude",
        "Historical data-fix audit. Do not ingest as a domain table.",
    ),
    "cleanup_audit_cancelled_trucking": (
        "Ops cleanup",
        "one-off cleanup residue",
        "EXCLUDE from analytics",
        "Exclude",
        "Historical trucking cancel/dedupe audit. Do not ingest as a domain table.",
    ),
    "cleanup_audit_duplicate_trucking_op_id": (
        "Ops cleanup",
        "one-off cleanup residue",
        "EXCLUDE from analytics",
        "Exclude",
        "Historical duplicate trucking-op audit.",
    ),
    "cleanup_audit_duplicate_trucking_per_contract": (
        "Ops cleanup",
        "one-off cleanup residue",
        "EXCLUDE from analytics",
        "Exclude",
        "Historical duplicate trucking-per-contract audit.",
    ),
}

COLUMN_HINTS: dict[str, str] = {
    "id": "UUID primary key generated by KLIP (uuid-ossp).",
    "created_at": "Row insert timestamp (timestamp without time zone; interpret as Asia/Jakarta unless noted).",
    "updated_at": "Last update timestamp (timestamp without time zone; interpret as Asia/Jakarta unless noted).",
    "created_by": "FK to users.id of the creator.",
    "updated_by": "FK to users.id of the last editor.",
    "contract_id": "Business contract number and/or FK depending on table — see table. On contracts this is the business key; on child tables it is often UUID FK to contracts.id.",
    "po_number": "Purchase order number. Part of the SAP tri-key with contract_number and sto_number.",
    "sto_number": "STO number. Part of the SAP tri-key. May be blank on some LAND/WB rows.",
    "contract_number": "SAP contract number (commercial identity).",
    "status": "Lifecycle status for this entity (values vary by table).",
    "password_hash": "Bcrypt password hash. NEVER land in the data lake unmasked.",
    "email": "User or party email. Treat as PII.",
    "phone": "Phone number. Treat as PII.",
    "username": "Login username.",
    "full_name": "Person display name. Treat as PII.",
    "role": "Application role name.",
    "is_active": "Soft-active flag.",
    "data": "JSONB payload. For sap_processed_data: sections contract, shipment, trucking[], vessel, payment, quality[], raw.",
    "raw_data": "Original JSON from the SAP Excel row before normalization.",
    "quantity_ordered": "Ordered/contract quantity (numeric). UI often shows MT; confirm UOM column.",
    "unit": "Unit of measure (MT, KG, etc.). Storage UOM flag may exist separately.",
    "incoterm": "Trade term: FOB, CIF, CFR, FRC, LCO, MIX, …",
    "transport_mode": "SEA or LAND (and related logistics classification).",
    "plant_code": "SAP plant code (join master_plants).",
    "group_name": "Vendor/plant group (region grouping in UI).",
    "buyer": "Buyer / company name from SAP.",
    "supplier": "Supplier name from SAP.",
    "product": "Product (CPO, PK, etc.).",
    "vessel_name": "Vessel name (match via master_vessels / aliases).",
    "vessel_code": "Vessel code from SAP/Jovin.",
    "operation_id": "Business operation id (SEA or LAND), unique per execution record.",
    "shipment_id": "Business shipment id (SEA).",
    "outstanding_qty": "Remaining quantity; often computed, not a physical stock figure.",
    "sap_presence": "PRESENT vs WITHDRAWN relative to latest trusted SAP import.",
    "klip_edited_fields": "JSONB list of fields last edited in KLIP (provenance vs SAP).",
    "hub_oidc_sub": "OIDC subject from Hub SSO. Treat as identifier; do not publish.",
}

SENSITIVE_COLUMNS = {
    ("users", "password_hash"),
    ("users", "email"),
    ("users", "phone"),
    ("users", "full_name"),
    ("users", "hub_oidc_sub"),
    ("users", "username"),
}

SENSITIVE_NAME_RE = re.compile(
    r"(password|secret|token|hash|otp|ssn|nik)",
    re.I,
)


def read_csv(name: str) -> list[dict[str, str]]:
    path = CSV_DIR / name
    raw = path.read_bytes()
    if raw.startswith(b"\xff\xfe") or raw.startswith(b"\xfe\xff"):
        text = raw.decode("utf-16")
    else:
        text = raw.decode("utf-8-sig")
    return list(csv.DictReader(text.splitlines()))


def fmt_type(row: dict[str, str]) -> str:
    dt = (row.get("data_type") or "").strip()
    udt = (row.get("udt_name") or "").strip()
    clen = row.get("character_maximum_length") or ""
    prec = row.get("numeric_precision") or ""
    scale = row.get("numeric_scale") or ""
    if dt == "character varying" or udt == "varchar":
        return f"VARCHAR({clen})" if clen else "VARCHAR"
    if dt == "character":
        return f"CHAR({clen})" if clen else "CHAR"
    if dt == "numeric" or udt == "numeric":
        if prec and scale not in ("", None):
            return f"NUMERIC({prec},{scale})"
        if prec:
            return f"NUMERIC({prec})"
        return "NUMERIC"
    if dt == "timestamp without time zone":
        return "TIMESTAMP"
    if dt == "timestamp with time zone":
        return "TIMESTAMPTZ"
    if dt == "ARRAY":
        return f"{udt.upper()}[]" if udt else "ARRAY"
    mapping = {
        "uuid": "UUID",
        "jsonb": "JSONB",
        "json": "JSON",
        "text": "TEXT",
        "boolean": "BOOLEAN",
        "date": "DATE",
        "integer": "INTEGER",
        "bigint": "BIGINT",
        "smallint": "SMALLINT",
        "double precision": "DOUBLE PRECISION",
        "real": "REAL",
        "bytea": "BYTEA",
        "inet": "INET",
        "interval": "INTERVAL",
    }
    if dt in mapping:
        return mapping[dt]
    if udt in mapping:
        return mapping[udt]
    return (dt or udt or "").upper() or "UNKNOWN"


def describe_column(table: str, col: str, pg_comment: str) -> str:
    if pg_comment:
        return pg_comment.strip()
    key = (table, col)
    specific = {
        ("contracts", "contract_id"): "Business contract number (SAP contract no). Not the UUID PK.",
        ("contracts", "id"): "KLIP UUID primary key. Child tables (shipments, payments, trucking) FK here.",
        ("contracts", "po_number"): "PO number. Together with contract_number/sto this is the commercial tri-key.",
        ("contracts", "sto_number"): "Primary/display STO on the contract row. Additional STOs live in contract_stos.",
        ("contracts", "quantity_ordered"): "Contract quantity. Displayed as MT in UI; confirm unit.",
        ("contracts", "sap_presence"): "PRESENT = in latest SAP report. WITHDRAWN = missed 2+ trusted imports (cancelled/deleted in SAP).",
        ("shipments", "contract_id"): "FK UUID → contracts.id.",
        ("shipments", "shipment_id"): "Business shipment identifier shown in UI.",
        ("shipments", "quantity_shipped"): "Loaded quantity (typically kg in DB; UI converts to MT).",
        ("shipments", "quantity_delivered"): "Delivered quantity (kg in DB).",
        ("shipments", "quantity_delivered_klip"): "KLIP-owned delivered qty (not copied from planning/SAP blindly).",
        ("payments", "contract_id"): "FK UUID → contracts.id.",
        ("payments", "payment_due_date"): "Payment due date used for aging / overdue.",
        ("payments", "payment_date"): "Actual payoff/payment date when paid.",
        ("sap_processed_data", "contract_number"): "Indexed identity extracted from JSONB for joins.",
        ("sap_processed_data", "po_number"): "Indexed PO identity. Tri-key with contract_number + sto_number.",
        ("sap_processed_data", "sto_number"): "Indexed STO identity. May be null; then JSON raw STO fields are fallback.",
        ("sap_processed_data", "data"): "Normalized JSONB. Flatten contract/shipment/payment for the lake; explode trucking[] and quality[].",
        ("trucking_operations", "contract_id"): "FK UUID → contracts.id.",
        ("vessel_loading_ports", "shipment_id"): "FK to shipments (UUID or business id depending on column — see data type).",
        ("users", "password_hash"): "Bcrypt hash of local-login password. Exclude or hash-redact in the lake.",
    }
    if key in specific:
        return specific[key]
    if col in COLUMN_HINTS:
        return COLUMN_HINTS[col]
    name = col.replace("_", " ")
    if col.endswith("_id") and col != "id":
        return f"Identifier/reference ({name}). Likely a UUID FK or business key — confirm on Relationships sheet."
    if col.endswith("_at"):
        return f"Timestamp for {name}."
    if col.endswith("_date"):
        return f"Date for {name}."
    if "qty" in col or "quantity" in col:
        return f"Quantity field ({name}). KLIP stores many logistics quantities in kg; UI shows MT (÷1000)."
    if col.endswith("_kg"):
        return f"Quantity in kilograms ({name}). Divide by 1000 for MT."
    if "status" in col:
        return f"Status flag ({name})."
    if col.startswith("eta_") or col.startswith("ata_") or col.startswith("etc_") or col.startswith("atc"):
        return f"SEA milestone date ({name}). ETA = estimate, ATA = actual. Port-level copy may exist on vessel_loading_ports."
    return f"{name[0].upper() + name[1:] if name else col} attribute on {table}."


def is_sensitive(table: str, col: str) -> str:
    if (table, col) in SENSITIVE_COLUMNS:
        return "YES — mask/exclude"
    if SENSITIVE_NAME_RE.search(col):
        return "YES — review"
    if table == "users" and col in {"email", "phone", "full_name", "username"}:
        return "YES — PII"
    return "No"


def style_header(ws: Worksheet, headers: list[str]) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    for col, h in enumerate(headers, 1):
        cell = ws.cell(1, col, h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def autosize(ws: Worksheet, widths: list[int]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = CELL_FONT
            cell.alignment = WRAP
            cell.border = THIN


def write_sheet(wb: Workbook, title: str, headers: list[str], rows: list[list[object]], widths: list[int]) -> Worksheet:
    ws = wb.create_sheet(title)
    style_header(ws, headers)
    prio_col = headers.index("Ingest priority") + 1 if "Ingest priority" in headers else None
    sens_col = headers.index("Sensitive") + 1 if "Sensitive" in headers else None
    for r_i, row in enumerate(rows, 2):
        for c_i, val in enumerate(row, 1):
            cell = ws.cell(r_i, c_i, val)
            if prio_col and c_i == prio_col:
                v = str(val)
                if v == "P0":
                    cell.fill = P0
                elif v == "P1":
                    cell.fill = P1
                elif v == "P2":
                    cell.fill = P2
                elif v == "Exclude":
                    cell.fill = EX
            if sens_col and c_i == sens_col and str(val).startswith("YES"):
                cell.fill = SENS
    autosize(ws, widths)
    ws.row_dimensions[1].height = 22
    ws.sheet_properties.tabColor = "1F4E79"
    return ws


def main() -> None:
    tables = read_csv("tmp_dl_tables.csv")
    columns = read_csv("tmp_dl_columns.csv")
    keys = read_csv("tmp_dl_keys.csv")
    fks = read_csv("tmp_dl_fks.csv")
    indexes = read_csv("tmp_dl_indexes.csv")
    counts = {r["table_name"]: r["exact_count"] for r in read_csv("tmp_dl_exact_counts.csv")}
    json_contract = read_csv("tmp_dl_json_contract.csv")
    json_top = read_csv("tmp_dl_json_top.csv")
    json_sec = read_csv("tmp_dl_json_sections.csv")

    pk_map: dict[str, list[str]] = defaultdict(list)
    uq_map: dict[str, list[str]] = defaultdict(list)
    for r in keys:
        if r["constraint_type"] == "PRIMARY KEY":
            pk_map[r["table_name"]].append(r["column_name"])
        elif r["constraint_type"] == "UNIQUE":
            uq_map[r["table_name"]].append(r["column_name"])

    fk_from: dict[str, list[str]] = defaultdict(list)
    for r in fks:
        fk_from[r["from_table"]].append(f"{r['from_column']} → {r['to_table']}.{r['to_column']}")

    wb = Workbook()
    # ReadMe
    ws = wb.active
    ws.title = "00_ReadMe"
    readme = [
        ["KLIP — Data Lake POC catalog"],
        ["Generated from", f"Live PostgreSQL information_schema + pg_catalog (local KLIP DB) on {AS_OF}"],
        ["Application", "KPN Logistics Intelligence Platform (KLIP)"],
        ["OLTP engine", "PostgreSQL 14"],
        ["Production DB (POC target)", "Aliyun Apsara PostgreSQL pgm-d9jn3khh0b3907w4.pgsql.ap-southeast-5.rds.aliyuncs.com:5432"],
        ["App servers", "FE 172.28.80.50 / BE 172.28.80.51 (VPC); BE is the only host that should reach Apsara"],
        ["GitHub", "https://github.com/dwsitproject-hub/Logistic  branch main (prod) / SIT (staging)"],
        [],
        ["How to use this workbook"],
        ["00_ReadMe", "This page"],
        ["01_Data_Structure", "Domains, grains, recommended lake layers, keys, quantity/UOM conventions"],
        ["02_Tables", "Every public table + 1 materialized view, with priority for the POC"],
        ["03_Columns", "Every column: type, nullability, PK/FK, description"],
        ["04_Relationships", "Declared foreign keys"],
        ["05_Indexes", "B-tree/expression indexes (useful for CDC extract predicates)"],
        ["06_SAP_JSONB", "Flatten map for sap_processed_data.data (the nested SAP payload)"],
        ["07_Lake_POC_Guidance", "CDC vs snapshot, PII, what NOT to ingest"],
        [],
        ["Important caveats"],
        ["Row counts", "Exact counts are from the LOCAL/dev database used to generate this file, NOT production Apsara. Use production COUNT(*) during the POC."],
        ["Derived snapshots", "Tables named *_snapshot and pipeline_*_summary are rebuildable caches for UI speed. Prefer CDC of source tables and rebuild gold in the lake."],
        ["Quantities", "Many execution qty columns are stored in kilograms. The KLIP UI divides by 1000 for MT."],
        ["Tri-key", "SAP uniqueness is contract_number + po_number + sto_number. Duplicates in the lake usually mean a broken tri-key."],
        ["Secrets", "Do not copy .env, JWT, SMTP, or GitHub tokens into the lake. Mask users.password_hash."],
        ["Cleanup tables", "cleanup_audit_* are one-off repair logs. Exclude from gold."],
    ]
    ws["A1"].font = Font(name="Calibri", size=16, bold=True, color="1F4E79")
    for i, row in enumerate(readme, 1):
        for j, val in enumerate(row, 1):
            cell = ws.cell(i, j, val)
            cell.font = Font(name="Calibri", size=11, bold=(j == 1 and i > 1))
            cell.alignment = WRAP
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 110
    ws.row_dimensions[1].height = 24
    ws.sheet_properties.tabColor = "1F4E79"

    # 01 structure
    struct_rows = [
        ["Commercial identity", "Tri-key", "contract_number + po_number + sto_number", "SAP ingest and duplicate detection. contracts.sto_number is only the primary STO; extras are in contract_stos."],
        ["Commercial master", "PO grain", "contracts.id (UUID) / contracts.po_number", "KLIP contract list is PO-centric. Outstanding/delivered qty is computed from SAP STO aggregation + qty-move overlay, not a single stored OS column."],
        ["SEA execution", "Shipment", "shipments.id / shipments.shipment_id", "One SEA shipment per operation. Timeline exists at shipment level AND vessel_loading_ports (port level)."],
        ["SEA ports", "Port call", "vessel_loading_ports.id", "Use this for ETA/ATA per loading/discharge port. Missing port rows can happen if only shipments.port_of_* were set."],
        ["LAND execution", "Trucking operation", "trucking_operations.id / operation_id", "Typically one active op per contract. Daily/WB facts in trucking_daily_actuals."],
        ["Finance", "Payment line", "payments.id", "Due vs paid; aging from payment_due_date vs payoff/payment_date."],
        ["SAP bronze", "Import session", "sap_data_imports.id", "Each Excel load. Row errors do not abort the import (SAVEPOINT pattern)."],
        ["SAP bronze", "Raw row", "sap_raw_data.id", "Keep for lineage/replay. Very wide JSON."],
        ["SAP silver", "Normalized row", "sap_processed_data.id + tri-key columns", "data JSONB must be flattened/exploded in the lake (see sheet 06)."],
        ["UI snapshots", "Denormalized cache", "snapshot PK + meta.updated_at", "Optional in bronze for debugging UI, but gold should be rebuilt from P0 tables."],
        ["UOM", "kg vs MT", "*_kg columns and many quantity_* fields", "DB kg → lake should store both kg and mt (qty_kg/1000.0) with explicit UOM."],
        ["Time", "naive timestamps", "TIMESTAMP without time zone", "Application timezone Asia/Jakarta. Recommend storing UTC in the lake and documenting offset +7."],
        ["Soft delete / withdraw", "sap_presence", "contracts.sap_presence, sap_processed_data presence fields", "WITHDRAWN rows remain for history; exclude from open-outstanding KPIs."],
        ["KLIP vs SAP", "provenance", "klip_edited_fields JSONB", "Fields edited in KLIP must not be blindly overwritten by the next SAP import."],
        ["B2B", "origin vs child PO", "empty Contract Reff PO on origin", "Origin OS/delivery may SUM children (capped at origin contract qty) when parent qty is null/0."],
        ["Transport split", "SEA vs LAND", "contracts.transport_mode / logistics_classification / sea_land in SAP JSON", "SAP distribution service routes to shipments vs trucking_operations."],
    ]
    write_sheet(
        wb,
        "01_Data_Structure",
        ["Area", "Concept", "Grain / key", "Description for data-lake design"],
        struct_rows,
        [22, 28, 55, 90],
    )

    table_rows = []
    for t in tables:
        name = t["table_name"]
        meta = TABLE_META.get(
            name,
            ("Unclassified", "See columns", "REVIEW", "P2", t.get("table_comment") or "See application code / columns sheet."),
        )
        domain, grain, layer, prio, desc = meta
        if t.get("table_comment") and not TABLE_META.get(name):
            desc = t["table_comment"]
        elif t.get("table_comment") and "See application" not in desc:
            pass
        table_rows.append(
            [
                name,
                t["object_type"],
                domain,
                grain,
                layer,
                prio,
                int(counts.get(name, 0) or 0),
                ", ".join(pk_map.get(name, [])),
                "; ".join(fk_from.get(name, [])[:8]),
                desc,
            ]
        )
    write_sheet(
        wb,
        "02_Tables",
        [
            "Table",
            "Object type",
            "Domain",
            "Grain",
            "Recommended lake layer",
            "Ingest priority",
            "Local row count (not prod)",
            "Primary key",
            "Foreign keys (sample)",
            "High-level description",
        ],
        table_rows,
        [38, 18, 20, 42, 32, 16, 22, 28, 55, 88],
    )

    col_rows = []
    for c in columns:
        table = c["table_name"]
        col = c["column_name"]
        meta = TABLE_META.get(table, ("Unclassified", "", "REVIEW", "P2", ""))
        pk = "PK" if col in pk_map.get(table, []) else ""
        uq = "UNIQUE" if col in uq_map.get(table, []) else ""
        fk_hits = [x for x in fk_from.get(table, []) if x.startswith(col + " ")]
        col_rows.append(
            [
                table,
                int(c["ordinal_position"] or 0),
                col,
                fmt_type(c),
                c.get("udt_name") or "",
                c.get("is_nullable") or "",
                (c.get("column_default") or "")[:180],
                pk,
                uq,
                "; ".join(fk_hits),
                is_sensitive(table, col),
                meta[3],
                describe_column(table, col, c.get("column_comment") or ""),
            ]
        )
    write_sheet(
        wb,
        "03_Columns",
        [
            "Table",
            "Pos",
            "Column",
            "Data type",
            "UDT",
            "Nullable",
            "Default",
            "PK",
            "Unique",
            "FK",
            "Sensitive",
            "Ingest priority",
            "High-level description",
        ],
        col_rows,
        [36, 6, 36, 22, 14, 10, 36, 6, 10, 42, 18, 14, 88],
    )

    fk_rows = [
        [r["from_table"], r["from_column"], r["to_table"], r["to_column"], r["update_rule"], r["delete_rule"], r["constraint_name"]]
        for r in fks
    ]
    write_sheet(
        wb,
        "04_Relationships",
        ["From table", "From column", "To table", "To column", "ON UPDATE", "ON DELETE", "Constraint"],
        fk_rows,
        [40, 28, 32, 22, 14, 14, 48],
    )

    idx_rows = [[r["tablename"], r["indexname"], r["indexdef"]] for r in indexes]
    write_sheet(
        wb,
        "05_Indexes",
        ["Table", "Index name", "Index definition"],
        idx_rows,
        [36, 55, 120],
    )

    jsonb_rows = []
    json_desc = {
        "contract": "Commercial header inside SAP JSON (buyer, PO, STO, qty, dates, plant, GR status).",
        "shipment": "SEA fields from SAP (vessel, ports, ETA/ATA family, B/L and delivery/receive qty).",
        "trucking": "JSON array of trucking legs (explode to rows). Element 0 keys listed as trucking[0].",
        "vessel": "Vessel particulars (LOA, draft, hull) complementing shipment.vessel_*.",
        "payment": "DP/due/payoff dates and deviation days from SAP finance columns.",
        "quality": "JSON array of quality blocks (location + nested data). Explode to rows.",
        "raw": "Original Excel header→value map (stringly typed). Keep in bronze; do not use as gold.",
    }
    for r in json_top:
        jsonb_rows.append(
            [
                "top-level",
                r["json_key"],
                "JSONB object or array",
                int(r["rows_with_key"]),
                json_desc.get(r["json_key"], "Section of sap_processed_data.data"),
            ]
        )
    for r in json_contract:
        jsonb_rows.append(
            [
                "data.contract",
                r["json_key"],
                "string/number in JSON (store as TEXT then cast)",
                int(r["rows_with_key"]),
                describe_column("sap_processed_data", r["json_key"], ""),
            ]
        )
    for r in json_sec:
        jsonb_rows.append(
            [
                f"data.{r['section']}",
                r["json_key"],
                "JSON value (cast in silver)",
                int(r["n"]),
                describe_column("sap_processed_data", r["json_key"], ""),
            ]
        )
    write_sheet(
        wb,
        "06_SAP_JSONB",
        ["JSON path", "Attribute", "Suggested lake type", "Local rows with key (not prod)", "Description"],
        jsonb_rows,
        [28, 42, 40, 28, 88],
    )

    guide = [
        ["P0 CDC / batch extract", "sap_data_imports, sap_raw_data, sap_processed_data, contracts, contract_stos, shipments, vessel_loading_ports, trucking_operations, trucking_daily_actuals, payments, master_plants, master_vessels", "Minimum viable lake for logistics + finance facts."],
        ["P1 enrich", "quality_surveys, documents (metadata only), audit_logs, sap_presence_audit, sap_import_failures, claim_susut_rows, claim_mutu_rows, trucking_realizations, suppliers, products, master_loading_ports, master_vessel_code_aliases", "Quality, claims, audit, masters."],
        ["Rebuild in lake (do not CDC as SoT)", "*_snapshot, *_snapshot_meta, shipment_pipeline_*, trucking_pipeline_*, trucking_list_stage_snapshot, mv_contract_payment_dates", "These exist because UI queries on JSONB were too slow. Recompute gold from P0."],
        ["Flatten JSONB", "sap_processed_data.data", "Silver tables: spd_contract, spd_shipment, spd_payment, spd_trucking (explode array), spd_quality (explode array), spd_raw (bronze)."],
        ["Natural keys", "tri-key + operation_id + shipment_id", "Keep UUID PK as surrogate; publish business keys for joins to SAP/other lakes."],
        ["CDC watermark", "updated_at, created_at, last_seen_at", "Most tables have updated_at. sap_processed_data also has last_seen_at / consecutive_misses for absence tracking."],
        ["Deletes", "Rare physical deletes; WITHDRAWN + sap_presence", "Prefer status/presence over hard delete. cleanup_audit_* prove one-off physical cleanups happened historically."],
        ["PII / secrets", "users.password_hash, users.email/phone, OIDC sub", "Hash-redact or exclude. Documents are metadata; binary files stay on NAS/volume."],
        ["Exclude", "cleanup_audit_*, schema_migrations (except a schema-version control table), user_preferences", "Noise / one-off."],
        ["KPI examples", "OTIF / milestone attainment, outstanding qty accuracy, payment aging, SAP failed-row %", "Define each KPI against GOLD tables in the lake, not against UI snapshots."],
        ["Extract connectivity", "Production Apsara allow-list is BE private IP 172.28.80.51 today", "Data-lake extractor needs its own allow-listed IP or a read replica; do not expose 5432 to internet."],
        ["Refresh jobs in KLIP", "Snapshot + pipeline summary refresh on backend", "If the lake replaces snapshots, KLIP can keep using OLTP snapshots; lake is analytics-only for the POC."],
        ["Sample local volumes (dev DB)", f"sap_raw_data={counts.get('sap_raw_data')} sap_processed_data={counts.get('sap_processed_data')} contracts={counts.get('contracts')} shipments={counts.get('shipments')} trucking_operations={counts.get('trucking_operations')} payments={counts.get('payments')} quality_surveys={counts.get('quality_surveys')}", "Production will differ; re-count on Apsara."],
    ]
    write_sheet(
        wb,
        "07_Lake_POC_Guidance",
        ["Topic", "Tables / keys", "Guidance"],
        guide,
        [28, 70, 95],
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
